import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PRIMARY = RGBColor(15, 35, 95)
SECONDARY = RGBColor(27, 92, 161)
ACCENT = RGBColor(0, 173, 181)
LIGHT = RGBColor(245, 247, 252)
DARK = RGBColor(34, 40, 49)
WHITE = RGBColor(255, 255, 255)


class ProfileDeckBuilder:
    def __init__(self, config_path: Path, output_path: Path):
        self.config = json.loads(config_path.read_text(encoding="utf-8"))
        self.output_path = output_path
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def add_background(self, slide, color=LIGHT):
        bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    def add_title(self, slide, title, subtitle=None):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.8), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.clear()
        p = title_frame.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(30)
        p.font.color.rgb = PRIMARY

        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.8), Inches(0.5))
            sp = sub_box.text_frame.paragraphs[0]
            sp.text = subtitle
            sp.font.size = Pt(15)
            sp.font.color.rgb = SECONDARY

    def add_card(self, slide, left, top, width, height, heading, lines):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RGBColor(220, 226, 240)
        shape.line.width = Pt(1)

        header = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.4))
        hp = header.text_frame.paragraphs[0]
        hp.text = heading
        hp.font.bold = True
        hp.font.size = Pt(15)
        hp.font.color.rgb = PRIMARY

        body = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.6), Inches(width - 0.4), Inches(height - 0.75))
        tf = body.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {line}"
            p.level = 0
            p.font.size = Pt(12)
            p.font.color.rgb = DARK
            p.space_after = Pt(8)

    def add_badges(self, slide, items, left, top, width, per_row=3, style="tech"):
        gap_x = 0.2
        gap_y = 0.2
        item_width = (width - gap_x * (per_row - 1)) / per_row
        item_height = 0.55

        for idx, item in enumerate(items):
            row = idx // per_row
            col = idx % per_row
            x = left + col * (item_width + gap_x)
            y = top + row * (item_height + gap_y)
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(item_width), Inches(item_height))
            shape.fill.solid()
            shape.line.fill.background()
            shape.fill.fore_color.rgb = RGBColor(233, 246, 255) if style == "tech" else RGBColor(233, 250, 242)
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = item
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = SECONDARY if style == "tech" else RGBColor(22, 122, 95)

    def slide_1_cover(self):
        c = self.config["company"]
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide, WHITE)

        brand = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        brand.fill.solid()
        brand.fill.fore_color.rgb = PRIMARY
        brand.line.fill.background()

        title = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(10), Inches(0.7))
        tp = title.text_frame.paragraphs[0]
        tp.text = c["name"]
        tp.font.bold = True
        tp.font.size = Pt(28)
        tp.font.color.rgb = WHITE

        year = slide.shapes.add_textbox(Inches(11.2), Inches(0.35), Inches(1.8), Inches(0.4))
        yp = year.text_frame.paragraphs[0]
        yp.text = c["year"]
        yp.alignment = PP_ALIGN.RIGHT
        yp.font.bold = True
        yp.font.size = Pt(18)
        yp.font.color.rgb = RGBColor(196, 228, 255)

        subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.8), Inches(0.7))
        sp = subtitle.text_frame.paragraphs[0]
        sp.text = c["subtitle"]
        sp.font.bold = True
        sp.font.size = Pt(24)
        sp.font.color.rgb = PRIMARY

        self.add_card(slide, 0.8, 2.9, 11.8, 2.6, "Executive Summary", [c["hybrid_advantage"]])

        footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.8), Inches(0.35))
        fp = footer.text_frame.paragraphs[0]
        fp.text = "Hybrid Advantage: UK security leadership + global delivery efficiency"
        fp.alignment = PP_ALIGN.CENTER
        fp.font.bold = True
        fp.font.size = Pt(13)
        fp.font.color.rgb = ACCENT

    def slide_2_capabilities(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "Core Capabilities & Technology Stack")
        self.add_card(slide, 0.8, 1.8, 6.0, 4.8, "Core Capabilities", self.config["capabilities"])
        self.add_card(slide, 7.0, 1.8, 5.5, 2.2, "Technology Stack", [])
        self.add_badges(slide, self.config["tech_stack"], left=7.25, top=2.45, width=5.0, per_row=2, style="tech")

    def slide_3_delivery_bridge(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "The Delivery Bridge & Governance")
        self.add_card(slide, 0.8, 1.8, 6.2, 4.8, "Delivery Bridge", self.config["delivery_bridge"])
        self.add_card(slide, 7.2, 1.8, 5.3, 4.8, "Data Sovereignty", self.config["data_sovereignty"])

    def slide_4_team_case_studies(self):
        c = self.config["company"]
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "Team Expertise, Case Studies & Corporate Info")

        studies = [f"{x['title']}: {x['outcome']}" for x in self.config["case_studies"][:3]]
        people = [f"{x['name']} - {x['bio']}" for x in self.config["personnel"]]
        corporate = [
            f"Address: {c['address']}",
            f"Companies House: {c['companies_house']}",
            f"ICO Registration: {c['ico_registration']}",
            f"Contact: {c['email']} | {c['phone']}"
        ]

        self.add_card(slide, 0.8, 1.8, 4.0, 4.8, "Case Studies", studies)
        self.add_card(slide, 4.95, 1.8, 4.0, 4.8, "Key Personnel", people)
        self.add_card(slide, 9.1, 1.8, 3.4, 4.8, "Corporate Info", corporate)

    def slide_5_certifications(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "Certifications & Accreditations")

        statement = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.8), Inches(0.7))
        sp = statement.text_frame.paragraphs[0]
        sp.text = "Security-first engineering with auditable compliance practices"
        sp.font.size = Pt(18)
        sp.font.bold = True
        sp.font.color.rgb = PRIMARY

        self.add_badges(slide, self.config["certifications"], left=1.0, top=2.9, width=11.3, per_row=2, style="compliance")

    def slide_6_market_positioning(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "Market Positioning")
        self.add_card(slide, 0.8, 1.8, 11.8, 4.8, "Differentiators & Value Proposition", self.config["market_positioning"])

    def slide_7_sla(self):
        sla = self.config["sla"]
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide)
        self.add_title(slide, "Service Level Agreement")
        self.add_card(slide, 0.8, 1.8, 5.7, 4.8, "Commitment Standards", [sla["uptime"], sla["response"], sla["resolution"], sla["reporting"]])
        self.add_card(slide, 6.7, 1.8, 5.9, 4.8, "Guarantees", [
            "Defined support windows and on-call escalation",
            "Transparent incident communications",
            "Service credits aligned to contractual KPIs",
            "Continuous improvement and root-cause actions"
        ])

    def slide_8_cta(self):
        c = self.config["company"]
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_background(slide, RGBColor(236, 245, 255))
        self.add_title(slide, "Call to Action & Contact", "Let's build secure, scalable digital services together.")

        hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.8), Inches(2.0))
        hero.fill.solid()
        hero.fill.fore_color.rgb = PRIMARY
        hero.line.fill.background()

        hp = hero.text_frame.paragraphs[0]
        hp.text = self.config["cta"]
        hp.alignment = PP_ALIGN.CENTER
        hp.font.bold = True
        hp.font.size = Pt(18)
        hp.font.color.rgb = WHITE

        contact = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11.8), Inches(1.4))
        tf = contact.text_frame
        tf.clear()
        lines = [
            c["name"],
            f"Email: {c['email']}   |   Phone: {c['phone']}",
            f"Website: {c['website']}"
        ]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = i == 0
            p.font.size = Pt(18 if i == 0 else 14)
            p.font.color.rgb = PRIMARY

    def build(self):
        self.slide_1_cover()
        self.slide_2_capabilities()
        self.slide_3_delivery_bridge()
        self.slide_4_team_case_studies()
        self.slide_5_certifications()
        self.slide_6_market_positioning()
        self.slide_7_sla()
        self.slide_8_cta()
        self.prs.save(str(self.output_path))


def main():
    base_dir = Path(__file__).resolve().parent
    config_path = base_dir / "config.json"
    output_path = base_dir / "WebFabricant_Company_Profile_2025.pptx"
    ProfileDeckBuilder(config_path, output_path).build()
    print(f"Presentation generated: {output_path}")


if __name__ == "__main__":
    main()
